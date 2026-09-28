import base64
import io
import logging
from collections import defaultdict
from datetime import datetime

from odoo import fields, models, _
from odoo.exceptions import UserError


_logger = logging.getLogger(__name__)


class ProductCatalogImportWizard(models.TransientModel):
    _name = "product.catalog.import.wizard"
    _description = "Product Catalogue Excel Import"

    file = fields.Binary(
        string="Excel File",
        required=True,
    )

    filename = fields.Char(
        string="Filename",
    )

    result_file = fields.Binary(
        string="Import Result",
        readonly=True,
    )

    result_filename = fields.Char(
        string="Result Filename",
        readonly=True,
    )

    summary = fields.Text(
        string="Import Summary",
        readonly=True,
    )

    COLUMN_FIELD_MAP = {
        "Name": "name",
        "Product Code": "default_code",
        "Film Type": "material_type",
        "Type": "type_reference",
        "Film Description": "film_description",
        "Is Metalize": "mo_serial_no",
        "Is Consumables": "is_consumables",
        "Treatment IN": "treatment_in_selection",
        "Treatment OUT": "treatment_out_selection",
        "Product Type": "type",
        "Unit of Measure": "uom_id",

        # Tracking is intentionally mapped here only so that
        # the Excel column is recognized.
        # Its value is NOT used.
        "Tracking": "tracking",

        "Control Policy": "purchase_method",
    }


    def action_import(self):

        self.ensure_one()

        if not self.file:
            raise UserError(
                _("Please upload an Excel file.")
            )

        filename = self.filename or ""

        if not filename.lower().endswith(".xlsx"):
            raise UserError(
                _("Only .xlsx Excel files are supported.")
            )

        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(
                _(
                    "Python package 'openpyxl' is not installed "
                    "on the Odoo server."
                )
            )
        try:
            file_data = base64.b64decode(self.file)

            workbook = load_workbook(
                filename=io.BytesIO(file_data),
                data_only=True,
            )

            sheet = workbook.active

        except Exception as e:
            raise UserError(
                _("Unable to read Excel file: %s") % str(e)
            )
        headers = []

        for cell in sheet[1]:

            value = self._clean_value(cell.value)

            if value:
                headers.append(value)
            else:
                headers.append("")

        required_columns = [
            "Product Code",
        ]

        missing_columns = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing_columns:
            raise UserError(
                _(
                    "Required column(s) missing from Excel:\n%s"
                )
                % ", ".join(missing_columns)
            )

        header_index = {
            header: index
            for index, header in enumerate(headers)
            if header
        }
        sku_rows = defaultdict(list)

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2),
            start=2,
        ):

            sku = self._get_cell_value(
                row,
                header_index,
                "Product Code",
            )

            sku = self._normalize_sku(sku)

            if sku:
                sku_rows[sku].append(row_number)

        duplicate_excel_skus = {
            sku: rows
            for sku, rows in sku_rows.items()
            if len(rows) > 1
        }
        results = []

        created_count = 0
        updated_count = 0
        error_count = 0
        skipped_count = 0

        ProductTemplate = self.env[
            "product.template"
        ].with_company(self.env.company)

        ProductProduct = self.env[
            "product.product"
        ].with_company(self.env.company)
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2),
            start=2,
        ):

            sku = self._get_cell_value(
                row,
                header_index,
                "Product Code",
            )

            sku = self._normalize_sku(sku)
            if not sku:

                skipped_count += 1

                results.append({
                    "row": row_number,
                    "sku": "",
                    "name": "",
                    "status": "Skipped",
                    "product_id": "",
                    "message": (
                        "Skipped because Product Code / SKU is blank"
                    ),
                })

                continue
            if sku in duplicate_excel_skus:

                error_count += 1

                results.append({
                    "row": row_number,
                    "sku": sku,
                    "name": "",
                    "status": "Error",
                    "product_id": "",
                    "message": (
                        "Duplicate SKU in Excel. "
                        "Rows: %s"
                        % ", ".join(
                            str(x)
                            for x in duplicate_excel_skus[sku]
                        )
                    ),
                })

                continue

            raw_values = {}

            try:
                with self.env.cr.savepoint():
                    for (
                        column_name,
                        field_name,
                    ) in self.COLUMN_FIELD_MAP.items():

                        if column_name not in header_index:
                            continue

                        value = self._get_cell_value(
                            row,
                            header_index,
                            column_name,
                        )

                        raw_values[field_name] = value
                    name = self._clean_value(
                        raw_values.get("name")
                    )

                    if not name:
                        name = sku
                    existing_products = ProductProduct.search(
                        [
                            (
                                "default_code",
                                "=ilike",
                                sku,
                            )
                        ]
                    )
                    if len(existing_products) > 1:

                        raise UserError(
                            _(
                                "SKU '%s' already exists on "
                                "multiple Odoo products: %s. "
                                "Import skipped to avoid updating "
                                "the wrong product."
                            )
                            % (
                                sku,
                                ", ".join(
                                    existing_products.mapped(
                                        lambda p: (
                                            "%s (ID %s)"
                                            % (
                                                p.display_name,
                                                p.id,
                                            )
                                        )
                                    )
                                ),
                            )
                        )
                    if existing_products:
                        product = existing_products[0]
                        template = product.product_tmpl_id

                        vals = self._prepare_product_values(
                            raw_values=raw_values,
                            sku=sku,
                            existing_product=product,
                            is_create=False,
                        )

                        if vals:
                            template.write(vals)
                        self._force_inventory_tracking(
                            template
                        )

                        updated_count += 1

                        results.append({
                            "row": row_number,
                            "sku": sku,
                            "name": name,
                            "status": "Updated",
                            "product_id": product.id,
                            "message": (
                                "Existing product updated. "
                                "Inventory tracking set to By Lots."
                            ),
                        })
                    else:
                        vals = self._prepare_product_values(
                            raw_values=raw_values,
                            sku=sku,
                            existing_product=False,
                            is_create=True,
                        )

                        vals["name"] = name
                        vals["is_storable"] = True
                        vals["tracking"] = "lot"

                        template = ProductTemplate.create(
                            vals
                        )
                        self._force_inventory_tracking(
                            template
                        )

                        product = template.product_variant_id

                        created_count += 1

                        results.append({
                            "row": row_number,
                            "sku": sku,
                            "name": name,
                            "status": "Created",
                            "product_id": product.id,
                            "message": (
                                "New product created. "
                                "Inventory tracking set to By Lots."
                            ),
                        })

            except Exception as e:

                error_count += 1

                error_message = self._format_error(
                    e
                )

                _logger.exception(
                    "Product catalogue import failed "
                    "for Excel row %s / SKU %s",
                    row_number,
                    sku,
                )

                results.append({
                    "row": row_number,
                    "sku": sku,
                    "name": (
                        self._clean_value(
                            raw_values.get("name")
                        )
                        or sku
                    ),
                    "status": "Error",
                    "product_id": "",
                    "message": error_message,
                })

                continue

        result_file = self._create_result_excel(
            original_file_data=file_data,
            results=results,
        )

        summary = (
            "Product Catalogue Import Completed\n\n"
            "Total Rows Processed: %s\n"
            "Created: %s\n"
            "Updated: %s\n"
            "Errors: %s\n"
            "Skipped: %s"
        ) % (
            len(results),
            created_count,
            updated_count,
            error_count,
            skipped_count,
        )

        self.write({
            "result_file": base64.b64encode(
                result_file
            ),
            "result_filename": (
                "product_catalog_import_result_%s.xlsx"
                % datetime.now().strftime(
                    "%Y%m%d_%H%M%S"
                )
            ),
            "summary": summary,
        })

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "view_mode": "form",
            "res_id": self.id,
            "target": "new",
        }


    def _force_inventory_tracking(self, product_template):

        ProductTemplate = self.env[
            "product.template"
        ]

        if "is_storable" in ProductTemplate._fields:

            product_template.write({
                "is_storable": True,
            })

        if "tracking" in ProductTemplate._fields:

            product_template.write({
                "tracking": "lot",
            })

        fields_to_flush = []

        if "is_storable" in ProductTemplate._fields:
            fields_to_flush.append("is_storable")

        if "tracking" in ProductTemplate._fields:
            fields_to_flush.append("tracking")

        if fields_to_flush:
            product_template.flush_recordset(
                fields_to_flush
            )
        product_template.invalidate_recordset(
            fields_to_flush
        )

        actual_tracking = (
            product_template.tracking
            if "tracking" in ProductTemplate._fields
            else False
        )

        actual_storable = (
            product_template.is_storable
            if "is_storable" in ProductTemplate._fields
            else False
        )

        if "is_storable" in ProductTemplate._fields:
            if not actual_storable:

                raise UserError(
                    _(
                        "Product '%s' could not be set as "
                        "Storable."
                    )
                    % product_template.display_name
                )

        if "tracking" in ProductTemplate._fields:
            if actual_tracking != "lot":

                raise UserError(
                    _(
                        "Product '%s' could not be set to "
                        "Lot tracking. Current value: %s"
                    )
                    % (
                        product_template.display_name,
                        actual_tracking,
                    )
                )

    def _prepare_product_values(
        self,
        raw_values,
        sku,
        existing_product=False,
        is_create=False,
    ):

        vals = {}

        ProductTemplate = self.env[
            "product.template"
        ]
        name = self._clean_value(
            raw_values.get("name")
        )

        if name:
            vals["name"] = name

        elif is_create:
            vals["name"] = sku
        vals["default_code"] = sku
        normal_fields = [
            "material_type",
            "type_reference",
            "film_description",
        ]

        for field_name in normal_fields:

            value = self._clean_value(
                raw_values.get(field_name)
            )

            if self._is_blank(value):
                continue

            if field_name not in ProductTemplate._fields:
                continue

            vals[field_name] = value
        boolean_fields = [
            "mo_serial_no",
            "is_consumables",
        ]

        for field_name in boolean_fields:

            raw_value = raw_values.get(field_name)

            if self._is_blank(raw_value):
                continue

            if field_name not in ProductTemplate._fields:
                continue

            vals[field_name] = self._parse_boolean(
                raw_value
            )
        product_type = self._clean_value(
            raw_values.get("type")
        )

        if not self._is_blank(product_type):

            field_name = "type"

            if field_name in ProductTemplate._fields:

                resolved_value = (
                    self._resolve_selection_value(
                        model=ProductTemplate,
                        field_name=field_name,
                        value=product_type,
                        aliases={
                            "storable": "product",
                            "storable product": "product",
                            "stockable": "product",
                            "consumable": "consu",
                            "consu": "consu",
                            "service": "service",
                        },
                    )
                )

                if resolved_value is not None:
                    vals[field_name] = resolved_value
        uom_name = self._clean_value(
            raw_values.get("uom_id")
        )

        if not self._is_blank(uom_name):

            field_name = "uom_id"

            if field_name in ProductTemplate._fields:

                uom = self.env["uom.uom"].search(
                    [
                        (
                            "name",
                            "=ilike",
                            uom_name,
                        )
                    ],
                    limit=1,
                )

                if uom:
                    vals["uom_id"] = uom.id

                else:
                    _logger.warning(
                        "UOM '%s' was not found. "
                        "Skipping UOM for SKU '%s'.",
                        uom_name,
                        sku,
                    )

        if "is_storable" in ProductTemplate._fields:
            vals["is_storable"] = True

        if "tracking" in ProductTemplate._fields:
            vals["tracking"] = "lot"

        treatment_in = self._clean_value(
            raw_values.get(
                "treatment_in_selection"
            )
        )

        if not self._is_blank(treatment_in):

            field_name = "treatment_in_selection"

            if field_name in ProductTemplate._fields:

                resolved_value = (
                    self._resolve_selection_value(
                        model=ProductTemplate,
                        field_name=field_name,
                        value=treatment_in,
                    )
                )

                if resolved_value is not None:
                    vals[field_name] = resolved_value
        treatment_out = self._clean_value(
            raw_values.get(
                "treatment_out_selection"
            )
        )

        if not self._is_blank(treatment_out):

            field_name = "treatment_out_selection"

            if field_name in ProductTemplate._fields:

                resolved_value = (
                    self._resolve_selection_value(
                        model=ProductTemplate,
                        field_name=field_name,
                        value=treatment_out,
                    )
                )

                if resolved_value is not None:
                    vals[field_name] = resolved_value

        control_policy = self._clean_value(
            raw_values.get("purchase_method")
        )

        if not self._is_blank(control_policy):

            field_name = "purchase_method"

            if field_name in ProductTemplate._fields:

                resolved_value = (
                    self._resolve_selection_value(
                        model=ProductTemplate,
                        field_name=field_name,
                        value=control_policy,
                        aliases={
                            "purchase": "purchase",
                            "ordered": "purchase",
                            "ordered quantities": "purchase",

                            "receive": "receive",
                            "received": "receive",
                            "received quantities": "receive",
                        },
                    )
                )

                if resolved_value is not None:
                    vals[field_name] = resolved_value

        return vals

    def _resolve_selection_value(
        self,
        model,
        field_name,
        value,
        aliases=None,
    ):

        field = model._fields.get(field_name)

        if not field:
            return None

        if self._is_blank(value):
            return None

        value_string = str(value).strip()

        normalized = (
            value_string
            .lower()
            .replace("_", " ")
            .replace("-", " ")
            .strip()
        )
        if aliases:
            alias_value = aliases.get(
                normalized
            )
            if alias_value:
                value_string = alias_value
        selection = field.get_description(
            self.env
        ).get("selection")

        if callable(selection):
            selection = selection(model)

        selection = selection or []
        for key, label in selection:
            if (
                str(key).strip().lower()
                == str(value_string).strip().lower()
            ):
                return key


        for key, label in selection:

            if (
                str(label).strip().lower()
                == str(value_string).strip().lower()
            ):
                return key
        _logger.warning(
            "Skipping invalid selection value '%s' "
            "for field '%s'. Allowed values: %s",
            value,
            field_name,
            ", ".join(
                "%s (%s)" % (
                    key,
                    label,
                )
                for key, label in selection
            ),
        )

        return None


    def _parse_boolean(self, value):

        if isinstance(value, bool):
            return value

        value = str(value).strip().lower()

        true_values = {
            "true",
            "yes",
            "y",
            "1",
            "t",
        }

        false_values = {
            "false",
            "no",
            "n",
            "0",
            "f",
        }

        if value in true_values:
            return True

        if value in false_values:
            return False

        raise UserError(
            _(
                "Invalid boolean value '%s'. "
                "Use TRUE/FALSE."
            )
            % value
        )


    def _validate_field_exists(self, field_name):
        ProductTemplate = self.env[
            "product.template"
        ]

        if field_name not in ProductTemplate._fields:

            raise UserError(
                _(
                    "Configured field '%s' "
                    "does not exist on product.template."
                )
                % field_name
            )

    def _create_result_excel(
        self,
        original_file_data,
        results,
    ):

        from openpyxl import load_workbook

        workbook = load_workbook(
            filename=io.BytesIO(
                original_file_data
            )
        )

        sheet = workbook.active
        headers = {
            str(cell.value).strip(): cell.column
            for cell in sheet[1]
            if cell.value
        }

        result_columns = [
            "Import Status",
            "Import Message",
            "Odoo Product ID",
            "Processed At",
        ]

        start_column = (
            sheet.max_column + 1
        )

        for index, column_name in enumerate(
            result_columns
        ):

            column_number = (
                start_column + index
            )

            if column_name in headers:

                column_number = headers[
                    column_name
                ]

            sheet.cell(
                row=1,
                column=column_number,
                value=column_name,
            )

        status_column = start_column
        message_column = start_column + 1
        product_id_column = start_column + 2
        processed_column = start_column + 3
        for result in results:

            row_number = result["row"]

            sheet.cell(
                row=row_number,
                column=status_column,
                value=result["status"],
            )

            sheet.cell(
                row=row_number,
                column=message_column,
                value=result["message"],
            )

            sheet.cell(
                row=row_number,
                column=product_id_column,
                value=result["product_id"],
            )

            sheet.cell(
                row=row_number,
                column=processed_column,
                value=datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            )
        widths = {
            status_column: 18,
            message_column: 70,
            product_id_column: 18,
            processed_column: 25,
        }

        for column_number, width in widths.items():

            sheet.column_dimensions[
                self._excel_column_letter(
                    column_number
                )
            ].width = width

        output = io.BytesIO()

        workbook.save(output)

        output.seek(0)

        return output.read()

    def _get_cell_value(
        self,
        row,
        header_index,
        column_name,
    ):

        if column_name not in header_index:
            return ""

        index = header_index[
            column_name
        ]

        if index >= len(row):
            return ""

        return self._clean_value(
            row[index].value
        )

    def _clean_value(self, value):

        if value is None:
            return ""

        if isinstance(value, str):
            return value.strip()

        return value

    def _is_blank(self, value):

        if value is None:
            return True

        if isinstance(value, str):
            return not value.strip()

        return False

    def _normalize_sku(self, sku):

        if sku is None:
            return ""

        return str(sku).strip()

    def _format_error(self, error):

        message = str(error)

        if not message:
            message = (
                error.__class__.__name__
            )

        return message

    def _excel_column_letter(self, column_number):

        result = ""

        while column_number:

            column_number, remainder = divmod(
                column_number - 1,
                26,
            )

            result = (
                chr(65 + remainder)
                + result
            )

        return result